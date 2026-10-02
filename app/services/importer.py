"""Bulk import: pasted lists, CSV and Excel files."""
from __future__ import annotations

import csv
import io
import re
from datetime import datetime

from sqlalchemy.orm import Session

from ..carriers import CARRIER_LABELS, CARRIERS, guess_carrier
from ..models import Client, Parcel, utcnow

FIELD_ALIASES = {
    "tracking_number": {"tracking_number", "tracking", "tracking number", "awb", "waybill",
                        "number", "tracking_no", "trackingno", "no", "colis"},
    "carrier": {"carrier", "courier", "transporteur", "company"},
    "client": {"client", "client_name", "customer", "account", "nom"},
    "description": {"description", "content", "goods", "item", "produit", "details"},
    "client_reference": {"reference", "client_reference", "ref", "order", "order_ref",
                         "invoice", "po"},
    "origin": {"origin", "from", "sender", "expediteur", "origin_city"},
    "destination": {"destination", "to", "receiver", "destinataire", "destination_city"},
    "cost_amount": {"cost", "cost_amount", "price", "amount", "montant", "freight"},
    "cost_currency": {"currency", "cost_currency", "devise"},
    "weight_kg": {"weight", "weight_kg", "kg", "poids"},
    "notes": {"notes", "note", "remark", "remarks", "comment", "comments"},
    "barcode": {"barcode", "bar_code", "ean", "upc", "scan", "code_barre"},
    "storage_location": {"storage_location", "storage", "shelf", "bin", "rack", "emplacement", "etagere"},
    "eta": {"eta", "estimated_delivery", "delivery_date", "date_livraison"},
}

# Extra spellings seen in real client spreadsheets. Built on top of the carrier
# registry so a carrier can never be missing from the import mapping.
_CARRIER_EXTRA_ALIASES = {
    "dhl": {"dhl express", "dhl_express", "dhlexpress", "dhl express sa"},
    "fedex": {"federal express", "fed ex", "fdx", "fx"},
    "aramex": {"aramex express", "ramex", "aramex courier"},
    "ups": {"united parcel service", "ups express", "ups saver"},
    "dhlecom": {"dhl ecommerce", "dhl e-commerce", "dhl parcel", "dhl global mail",
                "dhl ecommerce asia", "dhlecommerce"},
    "colissimo": {"colissimo france", "chronopost", "la poste", "laposte", "colis poste"},
    "track17": {"17 track", "17track", "ems", "post", "postal", "usps", "china post",
                "singapore post", "other carrier", "inconnu"},
    "manual": {"manual", "other", "autre", "unknown"},
}

CARRIER_ALIASES = {
    name: {name, name.replace("_", " "), CARRIER_LABELS[name].lower(),
           CARRIER_LABELS[name].split(" /")[0].lower(), CARRIER_LABELS[name].split(" (")[0].lower()}
          | _CARRIER_EXTRA_ALIASES.get(name, set())
    for name in CARRIERS
}


def _norm_key(key: str) -> str | None:
    k = re.sub(r"[\s_-]+", " ", (key or "").strip().lower())
    for field, aliases in FIELD_ALIASES.items():
        if k in {a.replace("_", " ") for a in aliases} or k.replace(" ", "_") in aliases:
            return field
    return None


def _norm_carrier(value: str | None, tracking_number: str) -> str:
    if value:
        v = re.sub(r"[\s_-]+", " ", str(value).strip().lower())
        if v:
            # exact match first, then longest-alias-wins so "dhl ecommerce" cannot be
            # swallowed by the shorter "dhl" entry
            for carrier, aliases in CARRIER_ALIASES.items():
                if v in aliases or v.replace(" ", "") in {a.replace(" ", "") for a in aliases}:
                    return carrier
            candidates = [(alias, carrier) for carrier, aliases in CARRIER_ALIASES.items()
                          for alias in aliases if alias and alias in v]
            if candidates:
                return max(candidates, key=lambda item: len(item[0]))[1]
    return guess_carrier(tracking_number)


def _parse_float(value) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "").replace(" ", ""))
    except ValueError:
        return None


def _parse_date(value) -> datetime | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def rows_to_parcels(db: Session, rows: list[dict], created_by_id: int | None = None,
                    auto_create_clients: bool = True) -> dict:
    """Turn raw dict/line rows into Parcel records (deduplicating by tracking number)."""
    created, skipped, errors = [], [], []
    client_cache: dict[str, Client] = {}

    for idx, raw in enumerate(rows, start=1):
        if isinstance(raw, str):
            parts = [p.strip() for p in raw.split(",")]
            raw = {"tracking_number": parts[0] if parts else ""}
            if len(parts) > 1 and parts[1]:
                raw["carrier"] = parts[1]
            if len(parts) > 2 and parts[2]:
                raw["client"] = parts[2]
            if len(parts) > 3 and parts[3]:
                raw["description"] = ", ".join(parts[3:])

        # normalise keys
        data: dict[str, object] = {}
        for key, value in raw.items():
            field = _norm_key(str(key))
            if field and value not in (None, ""):
                data[field] = value

        tracking = str(data.get("tracking_number", "")).strip()
        if not tracking:
            errors.append(f"line {idx}: no tracking number")
            continue

        existing = db.query(Parcel.id).filter(Parcel.tracking_number == tracking).first()
        if existing:
            skipped.append(tracking)
            continue

        carrier = _norm_carrier(str(data.get("carrier", "")) or None, tracking)
        client_id = None
        client_name = str(data.get("client", "")).strip()
        if client_name and auto_create_clients:
            client = client_cache.get(client_name.lower())
            if not client:
                client = db.query(Client).filter(Client.name.ilike(client_name)).first()
            if not client:
                client = Client(name=client_name)
                db.add(client)
                db.flush()
            client_cache[client_name.lower()] = client
            client_id = client.id

        parcel = Parcel(
            tracking_number=tracking,
            carrier=carrier,
            client_id=client_id,
            description=str(data.get("description", ""))[:2000],
            client_reference=str(data.get("client_reference", ""))[:120],
            origin=str(data.get("origin", ""))[:120],
            destination=str(data.get("destination", ""))[:120],
            cost_amount=_parse_float(data.get("cost_amount")),
            cost_currency=str(data.get("cost_currency", "MGA"))[:8],
            weight_kg=_parse_float(data.get("weight_kg")),
            notes=str(data.get("notes", ""))[:5000],
            barcode=str(data.get("barcode", ""))[:120],
            storage_location=str(data.get("storage_location", ""))[:120],
            cartons_expected=max(1, int(_parse_float(data.get("cartons_expected")) or 1)),
            eta=_parse_date(data.get("eta")),
            created_by_id=created_by_id,
            status="registered",
        )
        db.add(parcel)
        created.append(tracking)

    db.commit()
    return {"created": created, "skipped": skipped, "errors": errors}


def parse_text(text: str) -> list[dict]:
    """Paste-friendly parser: one parcel per line, optionally comma separated."""
    rows = []
    for line in (text or "").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        rows.append(line)
    return rows


def parse_csv_bytes(content: bytes) -> list[dict]:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    rows = [row for row in reader]
    # if the file has no recognisable header, treat every non-empty cell as a tracking number
    if rows and not any(_norm_key(k) for k in (rows[0].keys() if rows[0] else [])):
        flat = []
        for line in text.splitlines():
            for cell in re.split(r"[,;\t]", line):
                cell = cell.strip()
                if cell:
                    flat.append({"tracking_number": cell})
        return flat
    return rows


def parse_xlsx_bytes(content: bytes) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(c) if c is not None else "" for c in next(rows_iter)]
    except StopIteration:
        return []
    recognised = [h for h in header if _norm_key(h)]
    if not recognised:
        return [{"tracking_number": str(c).strip()} for row in ws.iter_rows(values_only=True)
                for c in row if c not in (None, "")]
    out = []
    for row in rows_iter:
        if not any(c not in (None, "") for c in row):
            continue
        out.append({header[i] if i < len(header) else f"col{i}": v for i, v in enumerate(row)})
    return out
