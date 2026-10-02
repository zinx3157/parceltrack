"""Warehouse receiving API: scan, carton check-in, stock, release."""
from __future__ import annotations

from fastapi import APIRouter, Depends, HTTPException, Response
from pydantic import BaseModel
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import Parcel, User
from ..security import current_user, require_operator
from ..services import receiving
from ..services.parcels import parcel_to_dict

router = APIRouter(prefix="/api/receiving", tags=["receiving"])


class ScanRequest(BaseModel):
    code: str
    storage_location: str = ""
    note: str = ""
    mark_received: bool = True


class UndoRequest(BaseModel):
    reason: str = ""


class ReleaseRequest(BaseModel):
    released_to: str = ""
    note: str = ""


class CartonRequest(BaseModel):
    storage_location: str = ""
    note: str = ""
    count: int = 1


@router.get("/lookup")
def lookup(code: str, _: User = Depends(current_user), db: Session = Depends(get_db)):
    return receiving.lookup(db, code)


@router.post("/scan")
def scan(payload: ScanRequest, user: User = Depends(require_operator),
         db: Session = Depends(get_db)):
    """Find the parcel by barcode and check in one carton."""
    result = receiving.lookup(db, payload.code)
    if not result.get("found"):
        return {"found": False, "code": payload.code,
                "message": "No parcel with this barcode — add it first, then scan again"}
    parcel = db.get(Parcel, result["id"])
    if parcel.released_at:
        raise HTTPException(409, f"Already released to {parcel.released_to or 'the consignee'} "
                                 f"— cancel the release first if this is a mistake")
    before = result["cartons_received"]
    state = receiving.receive_carton(db, parcel, user=user,
                                     storage_location=payload.storage_location,
                                     note=payload.note, count=1,
                                     mark_received=payload.mark_received)
    state["already_received"] = before > 0
    state["message"] = ("Receipt was already complete — carton added as an extra"
                        if state["cartons_received"] > state["cartons_expected"]
                        else state["message"])
    return state


@router.post("/{parcel_id}/receive")
def receive(parcel_id: int, payload: CartonRequest, user: User = Depends(require_operator),
            db: Session = Depends(get_db)):
    parcel = db.get(Parcel, parcel_id)
    if not parcel:
        raise HTTPException(404, "Parcel not found")
    return receiving.receive_carton(db, parcel, user=user,
                                    storage_location=payload.storage_location,
                                    note=payload.note, count=max(1, payload.count))


@router.post("/{parcel_id}/undo")
def undo(parcel_id: int, payload: UndoRequest, user: User = Depends(require_operator),
         db: Session = Depends(get_db)):
    parcel = db.get(Parcel, parcel_id)
    if not parcel:
        raise HTTPException(404, "Parcel not found")
    receiving.undo_receive(db, parcel)
    return parcel_to_dict(parcel, include_events=True)


@router.post("/{parcel_id}/release")
def release(parcel_id: int, payload: ReleaseRequest, user: User = Depends(require_operator),
            db: Session = Depends(get_db)):
    parcel = db.get(Parcel, parcel_id)
    if not parcel:
        raise HTTPException(404, "Parcel not found")
    if not parcel.received_at:
        raise HTTPException(400, "This parcel has not been received yet")
    receiving.release(db, parcel, released_to=payload.released_to, note=payload.note, user=user)
    state = receiving.parcel_receiving_state(db, parcel)
    state["message"] = f"Released to {parcel.released_to or 'the consignee'}"
    return state


@router.post("/{parcel_id}/release/undo")
def undo_release(parcel_id: int, _: User = Depends(require_operator),
                 db: Session = Depends(get_db)):
    parcel = db.get(Parcel, parcel_id)
    if not parcel:
        raise HTTPException(404, "Parcel not found")
    receiving.undo_release(db, parcel)
    state = receiving.parcel_receiving_state(db, parcel)
    state["message"] = "Release cancelled — back in stock"
    return state


@router.get("/recent")
def recent(limit: int = 20, _: User = Depends(current_user), db: Session = Depends(get_db)):
    return receiving.recently_received(db, limit=limit)


@router.get("/pending")
def pending(limit: int = 20, _: User = Depends(current_user), db: Session = Depends(get_db)):
    return receiving.pending_inbound(db, limit=limit)


@router.get("/stock")
def stock(q: str = "", location: str = "", include_released: bool = False, limit: int = 500,
          _: User = Depends(current_user), db: Session = Depends(get_db)):
    return {
        "summary": receiving.stock_summary(db),
        "items": receiving.stock_rows(db, query=q, location=location,
                                      include_released=include_released, limit=limit),
    }


@router.get("/stock.xlsx")
def stock_xlsx(q: str = "", location: str = "", include_released: bool = False,
               _: User = Depends(current_user), db: Session = Depends(get_db)):
    """Excel export of what is on the shelf (or the full movement list)."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = receiving.stock_rows(db, query=q, location=location,
                                include_released=include_released, limit=10000)
    columns = [("tracking_number", "Tracking number"), ("carrier", "Carrier"),
               ("client_name", "Client"), ("client_reference", "Reference"),
               ("description", "Description"), ("storage_location", "Location"),
               ("cartons_received", "Cartons in"), ("cartons_expected", "Cartons expected"),
               ("received_at", "Received at"), ("received_by", "Received by"),
               ("days_on_shelf", "Days on shelf"), ("status", "Carrier status"),
               ("released_at", "Released at"), ("released_to", "Released to"),
               ("barcode", "Barcode")]

    wb = Workbook()
    ws = wb.active
    ws.title = "Warehouse stock"
    head_fill = PatternFill("solid", fgColor="0E7490")
    for col, (_, label) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = head_fill
        cell.font = Font(color="FFFFFF", bold=True)
        ws.column_dimensions[get_column_letter(col)].width = max(13, min(30, len(label) + 6))
    for r, row in enumerate(rows, start=2):
        for c, (key, _) in enumerate(columns, start=1):
            ws.cell(row=r, column=c, value=row.get(key))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Summary")
    data = receiving.stock_summary(db)
    summary.append(["Warehouse stock", __import__("datetime").datetime.utcnow()
                    .strftime("%Y-%m-%d %H:%M") + " UTC"])
    summary.append([])
    for key in ("in_warehouse", "cartons_on_shelf", "stale", "released_total",
                "stale_after_days"):
        summary.append([key.replace("_", " ").title(), data[key]])
    summary.column_dimensions["A"].width = 24

    buf = io.BytesIO()
    wb.save(buf)
    name = f"warehouse_stock_{__import__('datetime').datetime.utcnow().strftime('%Y%m%d_%H%M')}.xlsx"
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
